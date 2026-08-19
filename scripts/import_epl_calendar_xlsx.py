#!/usr/bin/env python3
"""Import fixture calendars from top-5 leagues Excel into Supabase."""

from __future__ import annotations

import argparse
import sys
import urllib.parse
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "app"))

import openpyxl  # noqa: E402
from supabase_teams import SupabaseError, _request, fetch_teams  # noqa: E402

DEFAULT_SEASON = "2026-27"


@dataclass(frozen=True)
class LeagueImport:
    sheet: str
    league_id: str
    label: str


LEAGUE_IMPORTS: dict[str, LeagueImport] = {
    "epl": LeagueImport(
        sheet="АПЛ",
        league_id="2613ee27-7e8d-4d18-bd5e-e3f525121848",
        label="Premier League",
    ),
    "la_liga": LeagueImport(
        sheet="Ла Лига",
        league_id="8a33ea15-b5e2-4962-8196-a31c8b9fa8fe",
        label="La Liga",
    ),
    "bundesliga": LeagueImport(
        sheet="Бундеслига",
        league_id="0e928134-ae08-48fd-8d2c-3539a994b054",
        label="Bundesliga",
    ),
    "serie_a": LeagueImport(
        sheet="Серия A",
        league_id="c95e9c68-5679-41ca-ac92-b5b3975bfb02",
        label="Serie A",
    ),
    "ligue_1": LeagueImport(
        sheet="Лига 1",
        league_id="f89e6854-2836-4540-9d82-b3ff4019dc6a",
        label="Ligue 1",
    ),
}

DEFAULT_LEAGUE = "epl"


def read_sheet_fixtures(xlsx_path: Path, sheet: str) -> list[tuple[str, str, str]]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"Лист {sheet!r} не найден. Доступны: {wb.sheetnames}")
    ws = wb[sheet]
    out: list[tuple[str, str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 3:
            continue
        raw_date, home, away = row[0], row[1], row[2]
        if not raw_date or not home or not away:
            continue
        if isinstance(raw_date, datetime):
            date_str = raw_date.strftime("%Y-%m-%d")
        else:
            date_str = str(raw_date).strip()[:10]
        out.append((date_str, str(home).strip(), str(away).strip()))
    wb.close()
    if not out:
        raise ValueError(f"На листе {sheet!r} нет матчей")
    return out


def team_name_map(league_id: str) -> dict[str, int]:
    return {t.name_team: t.id for t in fetch_teams(league_id)}


def fetch_season_id(league_id: str, label: str) -> int | None:
    q = urllib.parse.urlencode(
        {
            "select": "id,label",
            "leagues_id": f"eq.{league_id}",
            "label": f"eq.{label}",
        }
    )
    rows = _request("GET", f"/season?{q}")
    if not rows:
        return None
    return int(rows[0]["id"])


def create_season(league_id: str, label: str) -> int:
    rows = _request(
        "POST",
        "/season?select=id,label",
        body={"leagues_id": league_id, "label": label},
    )
    return int(rows[0]["id"])


def count_season_matches(season_id: int) -> int:
    q = urllib.parse.urlencode({"select": "id", "season_id": f"eq.{season_id}"})
    rows = _request("GET", f"/matches?{q}")
    return len(rows) if isinstance(rows, list) else 0


def _match_payload(season_id: int, match_date: str, home_id: int, away_id: int) -> dict:
    return {
        "season_id": season_id,
        "match_date": match_date,
        "home_team_id": home_id,
        "away_team_id": away_id,
        "is_neutral": False,
        "match_weight": 1.0,
        "derby_weight": 1.0,
        "neutral_weight": 1.0,
        "active": True,
        "home_rotation_code": "none",
        "away_rotation_code": "none",
        "motivation": True,
    }


def insert_matches(
    season_id: int,
    fixtures: list[tuple[str, str, str]],
    team_map: dict[str, int],
    *,
    batch_size: int = 50,
) -> int:
    missing: set[str] = set()
    payloads: list[dict] = []
    for match_date, home, away in fixtures:
        if home not in team_map:
            missing.add(home)
        if away not in team_map:
            missing.add(away)
        if home in team_map and away in team_map:
            payloads.append(
                _match_payload(season_id, match_date, team_map[home], team_map[away])
            )
    if missing:
        raise ValueError(
            "Команды отсутствуют в Supabase (таблица team): "
            + ", ".join(sorted(missing))
        )

    inserted = 0
    for i in range(0, len(payloads), batch_size):
        batch = payloads[i : i + batch_size]
        _request("POST", "/matches?select=id", body=batch)
        inserted += len(batch)
    return inserted


def import_league(
    xlsx_path: Path,
    cfg: LeagueImport,
    *,
    season_label: str,
    replace: bool,
    skip_existing: bool,
    dry_run: bool,
) -> dict[str, object]:
    fixtures = read_sheet_fixtures(xlsx_path, cfg.sheet)
    team_map = team_name_map(cfg.league_id)
    teams_used = sorted({n for _, h, a in fixtures for n in (h, a)})
    unknown = [t for t in teams_used if t not in team_map]

    print(f"\n=== {cfg.label} ({cfg.sheet}) ===")
    print(f"Матчей: {len(fixtures)}, команд: {len(teams_used)}")
    if unknown:
        raise ValueError(f"{cfg.label}: не найдены в Supabase: {', '.join(unknown)}")

    existing_id = fetch_season_id(cfg.league_id, season_label)
    existing_count = count_season_matches(existing_id) if existing_id else 0
    if existing_id:
        print(f"Сезон {season_label}: season_id={existing_id}, матчей={existing_count}")
    else:
        print(f"Сезон {season_label} ещё не создан")

    if dry_run:
        print("Dry-run: запись не выполнялась")
        return {
            "league": cfg.label,
            "matches": len(fixtures),
            "season_id": existing_id,
            "imported": 0,
            "dry_run": True,
        }

    if existing_id and existing_count:
        if skip_existing and not replace:
            print(f"Пропуск: сезон уже заполнен ({existing_count} матчей)")
            return {
                "league": cfg.label,
                "matches": len(fixtures),
                "season_id": existing_id,
                "imported": 0,
                "skipped": True,
                "dry_run": dry_run,
            }
        if not replace:
            raise ValueError(
                f"{cfg.label}: сезон {season_label} уже содержит {existing_count} матчей. "
                "Используйте --replace."
            )
        _request("DELETE", f"/matches?season_id=eq.{existing_id}")
        print(f"Удалено матчей: {existing_count}")

    season_id = existing_id or create_season(cfg.league_id, season_label)
    if not existing_id:
        print(f"Создан season_id={season_id}")

    n = insert_matches(season_id, fixtures, team_map)
    final_count = count_season_matches(season_id)
    print(f"Импортировано: {n} (всего в сезоне: {final_count})")
    print(f"Первый: {fixtures[0][0]} {fixtures[0][1]} vs {fixtures[0][2]}")
    print(f"Последний: {fixtures[-1][0]} {fixtures[-1][1]} vs {fixtures[-1][2]}")
    return {
        "league": cfg.label,
        "matches": len(fixtures),
        "season_id": season_id,
        "imported": n,
        "dry_run": False,
    }


def main() -> None:
    league_keys = ", ".join(LEAGUE_IMPORTS)
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path, help="Excel с календарями топ-5 лиг")
    parser.add_argument(
        "--season-label",
        default=DEFAULT_SEASON,
        help=f"Метка сезона (default: {DEFAULT_SEASON})",
    )
    parser.add_argument(
        "--league",
        choices=sorted(LEAGUE_IMPORTS),
        default=DEFAULT_LEAGUE,
        help=f"Лига для импорта (default: {DEFAULT_LEAGUE})",
    )
    parser.add_argument(
        "--all",
        action="store_true",
        help="Импортировать все 5 лиг из файла",
    )
    parser.add_argument(
        "--replace",
        action="store_true",
        help="Удалить существующие матчи сезона перед импортом",
    )
    parser.add_argument(
        "--skip-existing",
        action="store_true",
        help="Пропустить лиги, у которых сезон уже содержит матчи",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Только проверка, без записи в Supabase",
    )
    args = parser.parse_args()

    if not args.xlsx.is_file():
        raise SystemExit(f"Файл не найден: {args.xlsx}")

    keys = sorted(LEAGUE_IMPORTS) if args.all else [args.league]
    print(f"Файл: {args.xlsx.name}")
    print(f"Сезон: {args.season_label}")
    print(f"Лиги: {', '.join(keys)}")

    results: list[dict[str, object]] = []
    errors: list[str] = []
    for key in keys:
        cfg = LEAGUE_IMPORTS[key]
        try:
            results.append(
                import_league(
                    args.xlsx,
                    cfg,
                    season_label=args.season_label,
                    replace=args.replace,
                    skip_existing=args.skip_existing,
                    dry_run=args.dry_run,
                )
            )
        except (ValueError, SupabaseError) as exc:
            body = getattr(exc, "body", "")[:200] if isinstance(exc, SupabaseError) else ""
            msg = f"{cfg.label}: {exc}" + (f" ({body})" if body else "")
            errors.append(msg)
            print(f"ОШИБКА: {msg}")

    print("\n--- Итог ---")
    for r in results:
        print(
            f"{r['league']}: {r['matches']} матчей, "
            f"season_id={r['season_id']}, imported={r['imported']}"
        )
    if errors:
        raise SystemExit(f"Ошибки ({len(errors)}): " + "; ".join(errors))


if __name__ == "__main__":
    main()
